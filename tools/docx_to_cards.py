"""勉強用単語帳_全124語.docx → cards.json（フェーズ1）

使い方: python tools/docx_to_cards.py  （python-docx が必要）
説明文・使いどころ・つながりは原文のまま写す。改変しない。
"""
import json
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "勉強用単語帳_全124語.docx"
OUT = ROOT / "cards.json"
ALIASES = ROOT / "aliases.json"

USAGE_PREFIX = "使いどころ："
LINKS_PREFIX = "つながり："

# 見出し行: 「用語　2024・2026年度・既出／注記」または「用語　※周辺語／注記」
HEAD_KISHUTSU = re.compile(r"^(?P<term>.+?)　(?P<years>\d{4}(?:・\d{4})*)年度・既出(?:／(?P<note>.+))?$")
HEAD_SHUHEN = re.compile(r"^(?P<term>.+?)　※周辺語(?:／(?P<note>.+))?$")


def parse_cell(paras):
    if len(paras) != 4:
        raise ValueError(f"段落数が4ではない: {paras[:1]}")
    head, body, usage, links = paras
    m = HEAD_KISHUTSU.match(head)
    if m:
        exam = {"years": [int(y) for y in m["years"].split("・")], "kind": "既出"}
    else:
        m = HEAD_SHUHEN.match(head)
        if not m:
            raise ValueError(f"見出し行を解釈できない: {head!r}")
        # 周辺語の注記にある年（制定年など）は years に入れない
        exam = {"years": [], "kind": "周辺語"}
    if m["note"]:
        exam["note"] = m["note"]
    if not usage.startswith(USAGE_PREFIX) or not links.startswith(LINKS_PREFIX):
        raise ValueError(f"使いどころ／つながりの行が見つからない: {head!r}")
    term = m["term"]
    return {
        "id": term,
        "term": term,
        "set": "bunka",
        "exam": exam,
        "body": body,
        "usage": usage[len(USAGE_PREFIX):],
        "links": links[len(LINKS_PREFIX):].split("、"),
        "status": "approved",
    }, body


def main():
    doc = docx.Document(str(SRC))
    cards, originals, seen = [], {}, set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:  # 結合セルは重複して現れるので用語で除去
                card, body = parse_cell([p.text for p in cell.paragraphs])
                if card["id"] in seen:
                    continue
                seen.add(card["id"])
                cards.append(card)
                originals[card["id"]] = body

    aliases = json.loads(ALIASES.read_text(encoding="utf-8"))
    ids = {c["id"] for c in cards}

    total = sum(len(c["links"]) for c in cards)
    direct = sum(l in ids for c in cards for l in c["links"])
    with_alias = sum((l in ids) or (aliases.get(l) in ids) for c in cards for l in c["links"])
    unresolved = sorted({l for c in cards for l in c["links"] if l not in ids and aliases.get(l) not in ids})
    body_ok = all(len(c["body"]) == len(originals[c["id"]]) and c["body"] == originals[c["id"]] for c in cards)

    with OUT.open("w", encoding="utf-8") as f:
        f.write("[\n" + ",\n".join(json.dumps(c, ensure_ascii=False) for c in cards) + "\n]\n")

    # 書き出したファイルを読み戻して、原文との一致を確かめる
    written = {c["id"]: c["body"] for c in json.loads(OUT.read_text(encoding="utf-8"))}
    body_ok = body_ok and all(written[k] == v and len(written[k]) == len(v) for k, v in originals.items())

    print(f"件数: {len(cards)}")
    print(f"説明文の文字数一致: {'OK' if body_ok else 'NG'}（{len(cards)}件）")
    print(f"つながり: 全{total}本 / 直接解決 {direct} / 別名込み {with_alias}")
    print(f"未解決の参照先: {'、'.join(unresolved)}")
    ok = len(cards) == 124 and body_ok and direct == 365 and with_alias == 367
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
